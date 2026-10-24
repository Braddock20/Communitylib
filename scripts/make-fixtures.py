#!/usr/bin/env python3
"""Generate the document fixtures used by the test suites.

Reproducible: every fixture is created from fixed text, so re-running yields
equivalent files. Run:  python3 scripts/make-fixtures.py [output_dir]

Requires: python-docx, openpyxl, python-pptx, reportlab, pikepdf, pypdf, Pillow.
Large/hostile inputs (zip bombs, oversized files) are NOT stored in git: tests build
them in memory. The malicious-archive fixtures below are tiny and inert.
"""
import io, os, sys, struct, zipfile, zlib, csv, textwrap
from pathlib import Path

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / 'test' / 'fixtures')
OUT.mkdir(parents=True, exist_ok=True)

DEJAVU = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
DEJAVU_BOLD = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'

PROSE = [
    "Electric charge is a conserved quantity. Like charges repel and unlike charges attract.",
    "Coulomb's law gives the force between two point charges in a vacuum.",
    "The electric field of a point charge points radially away from a positive charge.",
    "Gauss's law relates the electric flux through a closed surface to the enclosed charge.",
    "A capacitor stores energy in the electric field between its plates.",
    "Magnetic fields are produced by moving charges and by changing electric fields.",
]

def write(name, data):
    p = OUT / name
    p.write_bytes(data if isinstance(data, bytes) else data.encode('utf-8'))
    return p

# ------------------------------------------------------------------ text PDFs (reportlab)
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
pdfmetrics.registerFont(TTFont('DV', DEJAVU))
pdfmetrics.registerFont(TTFont('DVB', DEJAVU_BOLD))

def draw_page(c, title, lines, font='DV', size=12):
    w, h = A4
    c.setFont('DVB', 16); c.drawString(56, h - 72, title)
    c.setFont(font, size)
    y = h - 110
    for ln in lines:
        for part in textwrap.wrap(ln, 85) or ['']:
            c.drawString(56, y, part); y -= 18
        y -= 6

def make_text_pdf():
    buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=A4)
    c.setTitle('Physics Notes (text layer)'); c.setAuthor('Fixture Generator')
    for i in range(3):
        draw_page(c, f'Physics notes - page {i+1}', PROSE[i*2:i*2+2] + [f'Paragraph {i+1}: extractable text layer present.'])
        c.showPage()
    c.save(); write('text-3p.pdf', buf.getvalue())
make_text_pdf()

def make_equation_pdf():
    buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=A4)
    draw_page(c, 'Electromagnetism formulae', [
        'Gauss law relates flux to enclosed charge:',
        'Φ = ∮ E·dA = Q / ε₀',
        'The field of a point charge:',
        'E = k q / r²',
        'Faraday induction:',
        'ε = − dΦ/dt',
        'Chemical: 2H₂ + O₂ → 2H₂O',
    ])
    c.showPage(); c.save(); write('equations-1p.pdf', buf.getvalue())
make_equation_pdf()

def make_table_pdf():
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    st = getSampleStyleSheet()
    data = [['Quantity', 'Symbol', 'Unit'], ['Charge', 'q', 'coulomb (C)'], ['Potential', 'V', 'volt (V)'], ['Capacitance', 'C', 'farad (F)']]
    t = Table(data); t.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.5, colors.black)]))
    doc.build([Paragraph('Units and symbols', st['Heading1']), Spacer(1, 12), t, Spacer(1, 12), Paragraph('Table above lists common electrical quantities.', st['Normal'])])
    write('table-1p.pdf', buf.getvalue())
make_table_pdf()

def make_blank_middle_pdf():
    buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=A4)
    draw_page(c, 'Page one has text', [PROSE[0]]); c.showPage()
    c.showPage()  # genuinely blank page
    draw_page(c, 'Page three has text', [PROSE[1]]); c.showPage(); c.save()
    write('blank-middle-3p.pdf', buf.getvalue())
make_blank_middle_pdf()

# ------------------------------------------------------------------ scanned / mixed PDFs (image only)
from PIL import Image, ImageDraw, ImageFont
def render_text_image(lines, size=(1240, 1754), font_size=34, title=None):
    img = Image.new('L', size, 255); d = ImageDraw.Draw(img)
    f = ImageFont.truetype(DEJAVU, font_size); fb = ImageFont.truetype(DEJAVU_BOLD, font_size + 12)
    y = 120
    if title: d.text((100, y), title, font=fb, fill=0); y += 110
    for ln in lines:
        for part in textwrap.wrap(ln, 52) or ['']:
            d.text((100, y), part, font=f, fill=0); y += font_size + 22
        y += 20
    return img

def make_scanned_pdf():
    pages = [render_text_image(PROSE[i*2:i*2+2], title=f'Scanned notes {i+1}') for i in range(2)]
    pages[0].save(OUT / 'scanned-2p.pdf', 'PDF', resolution=150, save_all=True, append_images=pages[1:])
make_scanned_pdf()

def make_scanned_blank_first_pdf():
    # Mirrors the PHY121 observation: page 1 yields no text, later pages do.
    blank = Image.new('L', (1240, 1754), 255)
    pages = [blank, render_text_image(PROSE[2:4], title='Scanned notes 2'), render_text_image(PROSE[4:6], title='Scanned notes 3')]
    pages[0].save(OUT / 'scanned-blank-first-3p.pdf', 'PDF', resolution=150, save_all=True, append_images=pages[1:])
make_scanned_blank_first_pdf()

def make_mixed_pdf():
    # page 1: text layer; page 2: image only
    buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=A4)
    draw_page(c, 'Mixed document: text page', [PROSE[0], PROSE[1]]); c.showPage()
    img = render_text_image([PROSE[2], PROSE[3]], title='Mixed document: image page')
    ib = io.BytesIO(); img.save(ib, 'PNG'); ib.seek(0)
    from reportlab.lib.utils import ImageReader
    c.drawImage(ImageReader(ib), 0, 0, width=A4[0], height=A4[1]); c.showPage(); c.save()
    write('mixed-2p.pdf', buf.getvalue())
make_mixed_pdf()

def make_scanned_equation_pdf():
    img = render_text_image(['Gauss law:', 'Φ = ∮ E·dA = Q / ε₀', 'E = k q / r²', 'Plain sentence after the formulae.'], title='Scanned equations')
    img.save(OUT / 'scanned-equations-1p.pdf', 'PDF', resolution=150)
make_scanned_equation_pdf()

# ------------------------------------------------------------------ rotated, encrypted, corrupt PDFs
import pypdf, pikepdf
def make_rotated_pdf():
    r = pypdf.PdfReader(str(OUT / 'text-3p.pdf')); w = pypdf.PdfWriter()
    for i, p in enumerate(r.pages):
        if i == 1: p.rotate(90)
        w.add_page(p)
    with open(OUT / 'rotated-3p.pdf', 'wb') as f: w.write(f)
make_rotated_pdf()

def make_encrypted_pdf():
    with pikepdf.open(OUT / 'text-3p.pdf') as pdf:
        pdf.save(OUT / 'encrypted.pdf', encryption=pikepdf.Encryption(user='secret', owner='owner'))
make_encrypted_pdf()

def make_corrupt():
    data = (OUT / 'text-3p.pdf').read_bytes()
    write('corrupt-truncated.pdf', data[: len(data) // 3])
    write('not-really.pdf', b'This is plain text pretending to be a PDF.\n' * 5)
    write('fake.docx', b'%PDF-1.4\n% a PDF renamed to .docx\n')
    write('corrupt.docx', b'PK\x03\x04' + b'\x00' * 40)
make_corrupt()

# ------------------------------------------------------------------ images
def make_images():
    img = render_text_image(PROSE[:2], size=(900, 700), font_size=30, title='Image note')
    rgb = img.convert('RGB')
    rgb.save(OUT / 'note.png'); rgb.save(OUT / 'note.jpg', quality=88); rgb.save(OUT / 'note.webp', quality=88)
    rot = rgb.rotate(90, expand=True); rot.save(OUT / 'note-rotated.png')
    Image.new('RGB', (900, 700), 'white').save(OUT / 'blank.png')
make_images()

# ------------------------------------------------------------------ Office documents
def make_docx():
    import docx
    d = docx.Document()
    d.core_properties.title = 'Lab report'; d.core_properties.author = 'Fixture Generator'
    d.add_heading('Electricity Lab Report', level=1)
    d.add_paragraph('This report summarises the capacitor experiment.')
    d.add_heading('Method', level=2)
    for item in ['Charge the capacitor', 'Record the voltage', 'Discharge through a resistor']:
        d.add_paragraph(item, style='List Number')
    for item in ['Capacitor', 'Resistor', 'Voltmeter']:
        d.add_paragraph(item, style='List Bullet')
    t = d.add_table(rows=3, cols=3); t.style = 'Table Grid'
    rows = [['Time (s)', 'Voltage (V)', 'Note'], ['0', '5.0', 'start'], ['10', '1.8', 'decay']]
    for r, vals in enumerate(rows):
        for c_, v in enumerate(vals): t.cell(r, c_).text = v
    d.add_paragraph('Conclusion: the voltage decays exponentially.')
    d.save(OUT / 'report.docx')
make_docx()

def make_xlsx():
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Marks'
    ws.append(['Student', 'CAT1', 'CAT2']); ws.append(['Alice', 18, 20]); ws.append(['Bob', 15, 17]); ws['D1'] = 'Total'; ws['D2'] = '=B2+C2'
    ws2 = wb.create_sheet('Units'); ws2.append(['Code', 'Title']); ws2.append(['PHY121', 'Electricity, Magnetism, Optics and Modern Physics']); ws2.append(['MAT101', 'Calculus'])
    wb.create_sheet('Empty')
    wb.save(OUT / 'marks.xlsx')
make_xlsx()

def make_pptx():
    from pptx import Presentation
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1]); s1.shapes.title.text = 'Electric fields'; s1.placeholders[1].text = 'Field lines start on positive charges.\nThey end on negative charges.'
    s1.notes_slide.notes_text_frame.text = 'Remember to draw the diagram.'
    s2 = prs.slides.add_slide(prs.slide_layouts[1]); s2.shapes.title.text = "Gauss's law"; s2.placeholders[1].text = 'Flux equals enclosed charge over epsilon zero.'
    s3 = prs.slides.add_slide(prs.slide_layouts[5]); s3.shapes.title.text = 'Summary'
    prs.save(OUT / 'lecture.pptx')
make_pptx()

def make_csv():
    with open(OUT / 'quoted.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['code', 'title', 'notes'])
        w.writerow(['PHY121', 'Electricity, Magnetism', 'says "hello"\nsecond line'])
        w.writerow(['MAT101', 'Calculus', ''])
    write('table.tsv', 'a\tb\tc\n1\t2\t3\n"x\ty"\tz\tw\n')
make_csv()
write('notes.txt', '\n'.join(PROSE) + '\n')
write('notes.md', '# Heading\n\nSome *markdown* text.\n\n- item one\n- item two\n\n| a | b |\n|---|---|\n| 1 | 2 |\n')
write('page.html', '<!doctype html><html><head><title>Sample</title><style>p{color:red}</style><script>alert(1)</script></head><body><h1>Heading</h1><p>Hello <b>world</b>.</p><ul><li>one</li><li>two</li></ul><table><tr><th>a</th><th>b</th></tr><tr><td>1</td><td>2</td></tr></table><img src=x onerror=alert(1)></body></html>')

# ------------------------------------------------------------------ ZIP fixtures (tiny, inert)
def zip_bytes(entries, compression=zipfile.ZIP_DEFLATED):
    b = io.BytesIO()
    with zipfile.ZipFile(b, 'w', compression) as z:
        for name, data in entries:
            z.writestr(name, data)
    return b.getvalue()

write('archive-valid.zip', zip_bytes([('notes.txt', (OUT / 'notes.txt').read_bytes()), ('data/quoted.csv', (OUT / 'quoted.csv').read_bytes()), ('report.docx', (OUT / 'report.docx').read_bytes())]))
write('archive-traversal.zip', zip_bytes([('ok.txt', b'fine'), ('../../etc/evil.txt', b'bad'), ('/abs/evil.txt', b'bad'), ('dir\\..\\evil.txt', b'bad')]))
inner = zip_bytes([('inner.txt', b'hello from nested')])
write('archive-nested.zip', zip_bytes([('level1.zip', zip_bytes([('level2.zip', zip_bytes([('level3.zip', inner)]))]))]))
write('archive-many-entries.zip', zip_bytes([(f'f{i:04d}.txt', b'x') for i in range(700)]))
print('fixtures written to', OUT)
for p in sorted(OUT.iterdir()):
    print(f'  {p.name:32s} {p.stat().st_size:>8d} bytes')
